"""Build labs/pmo_tracker.xlsx: a synthetic Excel PMO tracker for Claude Code labs.

Sheets: Read Me | Milestones | Risks | Dashboard
Blue text = cells you type in. Black text = formulas (do not overwrite).
"""
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE / "pmo_tracker.xlsx"

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")
CALC = Font(name=FONT, color="000000")
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
RED = PatternFill("solid", fgColor="F4B6B6")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")
DATE_FMT = "DD-MMM-YYYY"

RISKS = [
    ("R-01", "TT-MR-Vaccine-SiteB", "Bioreactor IQ/OQ delay pushes engineering batch", "Equipment", 8, 7, 3, "Engineering", "Vendor FAT punch-list daily call; parallel OQ protocol review"),
    ("R-02", "TT-MR-Vaccine-SiteB", "Analytical method transfer fails acceptance criteria", "Analytical", 9, 4, 4, "QC", "Pre-transfer comparability run on 3 lots"),
    ("R-03", "TT-MR-Vaccine-SiteB", "Health authority asks for extra stability data", "Regulatory", 7, 5, 6, "RA", "Pre-submission meeting; start bridging stability now"),
    ("R-04", "TT-Rabies-CDMO", "Single-source raw material supply interruption", "Supply", 9, 5, 5, "Procurement", "Safety stock 6 months; accelerate 2nd source"),
    ("R-05", "TT-Rabies-CDMO", "Scale-down model not representative", "Process", 6, 3, 4, "MSAT", "Statistical equivalence (TOST) on key CQAs"),
    ("R-06", "TT-MR-Vaccine-SiteB", "Key MSAT SME attrition", "Resource", 5, 3, 2, "PMO", "Knowledge-transfer plan; backup SME named"),
]


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def rag_format(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Red"'], fill=RED))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Amber"'], fill=AMBER))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Green"'], fill=GREEN))


def main():
    df = pd.read_csv(HERE / "tt_milestones.csv", parse_dates=["baseline_date", "forecast_date"])
    wb = Workbook()

    # Read Me
    rm = wb.active
    rm.title = "Read Me"
    notes = [
        "PMO Tracker (SYNTHETIC DEMO DATA, not real company or GMP data)",
        "",
        "How to use:",
        "Blue text = inputs you type. Black text = formulas; do not overwrite.",
        "Milestones: edit Status, Baseline, Forecast, Critical Path. Slip and RAG update automatically.",
        "Risks: score Severity, Occurrence, Detection 1-10 (pFMEA style). RPN = S x O x D.",
        "Dashboard: read-only roll-up per project.",
        "",
        "RAG rule (same as .claude/skills/pmo-status-report):",
        "Complete = Green. Red = Delayed, or critical path slip > 30 days.",
        "Amber = At Risk, or slip 8-30 days. Otherwise Green.",
        "Risk action threshold: RPN > 100 = Mitigate (assumption; set your own in Dashboard!B12).",
    ]
    for i, t in enumerate(notes, start=1):
        rm.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=(i == 1))
    rm.column_dimensions["A"].width = 100

    # Milestones
    ms = wb.create_sheet("Milestones")
    cols = ["Project", "Workstream", "Milestone", "Owner", "Baseline Date", "Forecast Date",
            "Status", "Critical Path (Y/N)", "Slip (days)", "RAG"]
    header(ms, cols, [24, 13, 46, 14, 14, 14, 13, 11, 10, 9])
    for r, row in enumerate(df.itertuples(index=False), start=2):
        vals = [row.project, row.workstream, row.milestone, row.owner,
                row.baseline_date.date(), row.forecast_date.date(), row.status, row.critical_path]
        for c, v in enumerate(vals, start=1):
            cell = ms.cell(row=r, column=c, value=v)
            cell.font = INPUT
            if c in (5, 6):
                cell.number_format = DATE_FMT
        ms.cell(row=r, column=9, value=f"=F{r}-E{r}").font = CALC
        ms.cell(row=r, column=10, value=(
            f'=IF(G{r}="Complete","Green",IF(OR(G{r}="Delayed",AND(H{r}="Y",I{r}>30)),"Red",'
            f'IF(OR(G{r}="At Risk",AND(I{r}>=8,I{r}<=30)),"Amber","Green")))')).font = CALC
    last_ms = len(df) + 1
    dv = DataValidation(type="list", formula1='"Not Started,In Progress,At Risk,Delayed,Complete"')
    ms.add_data_validation(dv)
    dv.add(f"G2:G{last_ms + 50}")
    rag_format(ms, f"J2:J{last_ms}")

    # Risks
    rk = wb.create_sheet("Risks")
    cols = ["Risk ID", "Project", "Risk", "Category", "Severity (1-10)", "Occurrence (1-10)",
            "Detection (1-10)", "RPN", "Action", "Owner", "Mitigation"]
    header(rk, cols, [8, 24, 46, 12, 10, 11, 10, 8, 10, 13, 55])
    for r, risk in enumerate(RISKS, start=2):
        for c, v in enumerate(risk[:7], start=1):
            rk.cell(row=r, column=c, value=v).font = INPUT
        rk.cell(row=r, column=8, value=f"=E{r}*F{r}*G{r}").font = CALC
        rk.cell(row=r, column=9, value=f'=IF(H{r}>Dashboard!$B$12,"Mitigate","Monitor")').font = CALC
        rk.cell(row=r, column=10, value=risk[7]).font = INPUT
        rk.cell(row=r, column=11, value=risk[8]).font = INPUT
    last_rk = len(RISKS) + 1
    rk.conditional_formatting.add(f"I2:I{last_rk}", CellIsRule(operator="equal", formula=['"Mitigate"'], fill=RED))

    # Dashboard
    db = wb.create_sheet("Dashboard")
    cols = ["Project", "Red", "Amber", "Green", "Project RAG", "Risks to Mitigate", "Worst Slip (days)"]
    header(db, cols, [26, 8, 8, 8, 12, 14, 14])
    M = f"Milestones!$A$2:$A${last_ms}"
    J = f"Milestones!$J$2:$J${last_ms}"
    H = f"Milestones!$H$2:$H${last_ms}"
    I = f"Milestones!$I$2:$I${last_ms}"
    for r, proj in enumerate(sorted(df.project.unique()), start=2):
        db.cell(row=r, column=1, value=proj).font = CALC
        for c, rag in zip((2, 3, 4), ("Red", "Amber", "Green")):
            db.cell(row=r, column=c, value=f'=COUNTIFS({M},$A{r},{J},"{rag}")').font = CALC
        db.cell(row=r, column=5, value=(
            f'=IF(COUNTIFS({M},$A{r},{H},"Y",{J},"Red")>0,"Red",'
            f'IF(COUNTIFS({M},$A{r},{H},"Y",{J},"Amber")>0,"Amber","Green"))')).font = CALC
        db.cell(row=r, column=6, value=(
            f'=COUNTIFS(Risks!$B$2:$B${last_rk},$A{r},Risks!$I$2:$I${last_rk},"Mitigate")')).font = CALC
        db.cell(row=r, column=7, value=f"=_xlfn.MAXIFS({I},{M},$A{r})").font = CALC
    rag_format(db, "E2:E10")
    db.cell(row=11, column=1, value="Settings").font = Font(name=FONT, bold=True)
    db.cell(row=12, column=1, value="RPN action threshold").font = CALC
    th = db.cell(row=12, column=2, value=100)
    th.font, th.fill = INPUT, PatternFill("solid", fgColor="FFFF00")
    db.cell(row=12, column=3, value="Assumption: common pFMEA cut-off; adjust to your site's risk SOP").font = Font(name=FONT, italic=True)
    db.cell(row=14, column=1, value=f"Built {date.today():%d-%b-%Y} from labs/tt_milestones.csv").font = Font(name=FONT, italic=True)

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
